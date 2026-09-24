from __future__ import annotations

import csv
import hashlib
import io
import json
import shutil
import subprocess
from pathlib import Path
from uuid import uuid4

from fastapi import UploadFile
from openpyxl import load_workbook
from pypdf import PdfReader

from .models import DataSource, SourceConnectionRequest

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
SUPPORTED_SUFFIXES = {".csv", ".xlsx", ".pdf", ".docx", ".doc"}


class SourceIngestionError(ValueError):
    pass


def connected_source(request: SourceConnectionRequest) -> DataSource:
    return DataSource(
        source_id=f"connection-{uuid4()}",
        name=request.name,
        kind=request.kind,
        status="configured",
        metadata={
            key: value
            for key, value in {"endpoint": request.endpoint, "dataset": request.dataset}.items()
            if value
        },
    )


async def ingest_upload(upload: UploadFile, uploads_dir: Path) -> tuple[DataSource, list[dict]]:
    filename = Path(upload.filename or "upload").name
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise SourceIngestionError("Use CSV, XLSX, PDF, DOCX, or DOC files")
    content = await upload.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise SourceIngestionError("Files must be 25 MB or smaller")
    if not content:
        raise SourceIngestionError("The uploaded file is empty")

    source_id = f"file-{uuid4()}"
    source_dir = uploads_dir / source_id
    source_dir.mkdir(parents=True, exist_ok=False)
    raw_path = source_dir / filename
    raw_path.write_bytes(content)
    digest = hashlib.sha256(content).hexdigest()
    try:
        records = _extract_records(raw_path, suffix)
    except Exception:
        shutil.rmtree(source_dir, ignore_errors=True)
        raise

    source = DataSource(
        source_id=source_id,
        name=filename,
        kind=f"file{suffix}",
        record_count=len(records),
        metadata={
            "filename": filename,
            "content_type": upload.content_type or "application/octet-stream",
            "size_bytes": len(content),
            "sha256": digest,
            "raw_path": str(raw_path),
        },
    )
    return source, records


def _extract_records(path: Path, suffix: str) -> list[dict]:
    if suffix == ".csv":
        return _read_csv(path)
    if suffix == ".xlsx":
        return _read_xlsx(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    return _read_doc(path)


def _read_csv(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    if not text.strip():
        return []
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel
    return [dict(row) for row in csv.DictReader(io.StringIO(text), dialect=dialect)]


def _read_xlsx(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    records: list[dict] = []
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or f"column_{index + 1}") for index, value in enumerate(next(rows, []))]
        for row_index, row in enumerate(rows, start=2):
            records.append(
                {
                    "_sheet": sheet.title,
                    "_row": row_index,
                    **{header: value for header, value in zip(headers, row, strict=False)},
                }
            )
    return records


def _read_pdf(path: Path) -> list[dict]:
    reader = PdfReader(path)
    return [
        {"page": index, "text": page.extract_text() or ""}
        for index, page in enumerate(reader.pages, start=1)
    ]


def _read_docx(path: Path) -> list[dict]:
    from docx import Document

    document = Document(path)
    records = [
        {"block": index, "text": paragraph.text}
        for index, paragraph in enumerate(document.paragraphs, start=1)
        if paragraph.text.strip()
    ]
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            records.append(
                {
                    "table": table_index,
                    "row": row_index,
                    "cells": [cell.text for cell in row.cells],
                }
            )
    return records


def _read_doc(path: Path) -> list[dict]:
    if not shutil.which("antiword"):
        raise SourceIngestionError("DOC extraction is unavailable on this host; save as DOCX")
    result = subprocess.run(
        ["antiword", str(path)], capture_output=True, text=True, timeout=30, check=False
    )
    if result.returncode:
        raise SourceIngestionError("The DOC file could not be read")
    return [
        {"block": index, "text": block.strip()}
        for index, block in enumerate(result.stdout.split("\n\n"), start=1)
        if block.strip()
    ]


def write_manifest(source: DataSource, records: list[dict], source_dir: Path) -> None:
    (source_dir / "manifest.json").write_text(
        json.dumps(
            {
                "source": source.model_dump(mode="json"),
                "record_count": len(records),
                "format_version": 1,
            },
            indent=2,
            default=str,
        )
    )


def archive_source(
    source: DataSource, records: list[dict], source_dir: Path, bucket_name: str
) -> DataSource:
    from google.cloud import storage

    records_path = source_dir / "records.jsonl"
    records_path.write_text(
        "\n".join(json.dumps(record, default=str) for record in records), encoding="utf-8"
    )
    source.metadata["archive_uri"] = f"gs://{bucket_name}/sources/{source.source_id}/"
    write_manifest(source, records, source_dir)
    client = storage.Client()
    bucket = client.bucket(bucket_name)
    for path in source_dir.iterdir():
        blob = bucket.blob(f"sources/{source.source_id}/{path.name}")
        blob.metadata = {"source_id": source.source_id, "sha256": source.metadata["sha256"]}
        blob.upload_from_filename(path)
    return source


def restore_archived_sources(store, bucket_name: str) -> int:
    from google.cloud import storage

    restored = 0
    client = storage.Client()
    bucket = client.bucket(bucket_name)
    for blob in client.list_blobs(bucket, prefix="sources/"):
        if not blob.name.endswith("/manifest.json"):
            continue
        manifest = json.loads(blob.download_as_text())
        source = DataSource.model_validate(manifest["source"])
        if store.get_source(source.source_id):
            continue
        records_blob = bucket.blob(f"sources/{source.source_id}/records.jsonl")
        records = [
            json.loads(line)
            for line in records_blob.download_as_text().splitlines()
            if line.strip()
        ]
        store.save_source(source, records)
        restored += 1
    return restored


def verify_bigquery(request: SourceConnectionRequest, billing_project: str) -> dict:
    import re

    from google.cloud import bigquery

    project = request.project or ""
    dataset_id = request.dataset or ""
    if not project and "." in dataset_id:
        project, dataset_id = dataset_id.split(".", 1)
    if not re.fullmatch(r"[a-z][a-z0-9-]{4,61}[a-z0-9]", project):
        raise SourceIngestionError("Enter a valid Google Cloud project ID")
    if not re.fullmatch(r"[A-Za-z0-9_]{1,1024}", dataset_id):
        raise SourceIngestionError("Enter the dataset ID without a project prefix")
    client = bigquery.Client(project=billing_project or project)
    dataset = client.get_dataset(f"{project}.{dataset_id}", timeout=15)
    if dataset.location.lower() != request.location.lower():
        raise SourceIngestionError(f"This dataset is in {dataset.location}; update the location")
    tables = list(client.list_tables(dataset, max_results=100, timeout=15))
    records: list[dict] = []
    total = 0
    for table in tables:
        if table.table_type != "TABLE":
            continue
        detail = client.get_table(table.reference, timeout=15)
        total += detail.num_rows or 0
        # Bounded read checks data access, not only metadata permissions.
        if len(records) < 200:
            rows = client.list_rows(detail, max_results=min(20, 200 - len(records)), timeout=15)
            records.extend({"_table": table.table_id, **dict(row.items())} for row in rows)
    return {
        "project": project,
        "dataset": dataset_id,
        "location": dataset.location,
        "tables": [table.table_id for table in tables],
        "record_count": total,
        "records": records,
    }
